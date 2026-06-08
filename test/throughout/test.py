"""
5-batch full-flow test runner.

Usage:
  python test/throughout/test.py batch1
  python test/throughout/test.py batch2
  python test/throughout/test.py batch3
  python test/throughout/test.py batch4
  python test/throughout/test.py batch5
  python test/throughout/test.py all

Reports are written to:
  test/throughout/5batch_samples/report/
"""
from __future__ import annotations

import argparse
import json
import sys
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[2]
SRC = ROOT / "src"
SAMPLES_FILE = ROOT / "test" / "throughout" / "5batch_samples" / "5batch_samples.xlsx"
REPORT_DIR = ROOT / "test" / "throughout" / "5batch_samples" / "report"

sys.path.insert(0, str(SRC))
from entrance import entrance  # noqa: E402


def parse_json_cell(value: str) -> Any:
    if value == "-":
        return "-"
    return json.loads(value)


def load_samples() -> list[dict[str, Any]]:
    wb = load_workbook(SAMPLES_FILE, read_only=True, data_only=True)
    ws = wb["samples"]
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        messages, correct_label, batch = row
        rows.append({
            "messages": json.loads(messages),
            "correct_label": parse_json_cell(correct_label),
            "batch": str(batch),
        })
    return rows


def batch_filter(samples: list[dict[str, Any]], batch: str) -> list[dict[str, Any]]:
    if batch == "batch1":
        return [s for s in samples if s["batch"].startswith("1.")]
    if batch == "batch2":
        return [s for s in samples if s["batch"].startswith("2.")]
    if batch == "batch3":
        return [s for s in samples if s["batch"] == "3"]
    if batch == "batch4":
        return [s for s in samples if s["batch"] == "4"]
    if batch == "batch5":
        return [s for s in samples if s["batch"].startswith("5.")]
    raise ValueError(f"unknown batch: {batch}")


def run_one(messages: list[dict[str, str]], debug: bool, l0_threshold: float) -> dict[str, Any]:
    return entrance(messages, l0_threshold=l0_threshold, debug=debug)


def output_risk(result: dict[str, Any]) -> str | None:
    return result.get("risk_level")


def output_intent(result: dict[str, Any]) -> dict[str, str | None]:
    pi = result.get("data", {}).get("primary_intent", {}) or {}
    return {"l1": pi.get("l1"), "l2": pi.get("l2")}


def messages_text(messages: list[dict[str, str]]) -> str:
    return json.dumps(messages, ensure_ascii=False)


def result_text(result: dict[str, Any]) -> str:
    return json.dumps(result, ensure_ascii=False, indent=2)


def report_path(name: str) -> Path:
    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    return REPORT_DIR / f"{name}_{ts}.xlsx"


def write_batch1_report(samples: list[dict[str, Any]], debug: bool, l0_threshold: float) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "batch1"
    ws.append(["测试messages", "该样本运行时输出的全部控制信息json"])
    for i, s in enumerate(samples, start=1):
        print(f"  batch1 {i}/{len(samples)}")
        result = run_one(s["messages"], debug=debug, l0_threshold=l0_threshold)
        ws.append([messages_text(s["messages"]), result_text(result)])
    ws.column_dimensions["A"].width = 90
    ws.column_dimensions["B"].width = 120
    path = report_path("batch1_report")
    wb.save(path)
    return path


def write_risk_report(name: str, samples: list[dict[str, Any]], debug: bool, l0_threshold: float) -> Path:
    rows = []
    correct = 0
    for i, s in enumerate(samples, start=1):
        print(f"  {name} {i}/{len(samples)}")
        result = run_one(s["messages"], debug=debug, l0_threshold=l0_threshold)
        expected = s["correct_label"].get("risk_level")
        actual = output_risk(result)
        if actual == expected:
            correct += 1
        rows.append([messages_text(s["messages"]), json.dumps(s["correct_label"], ensure_ascii=False), actual])

    total = len(samples)
    acc = correct / total if total else 0.0

    wb = Workbook()
    ws = wb.active
    ws.title = name
    ws.append([f"total={total}", f"correct={correct}", f"acc={acc:.4f}"])
    ws.append(["messages", "正确标签", "输出标签（风险等级）"])
    for row in rows:
        ws.append(row)
    ws.column_dimensions["A"].width = 90
    ws.column_dimensions["B"].width = 35
    ws.column_dimensions["C"].width = 20
    path = report_path(f"{name}_report")
    wb.save(path)
    return path


def write_routing_report(name: str, samples: list[dict[str, Any]], debug: bool, l0_threshold: float) -> Path:
    rows = []
    exact = 0
    l1_correct = 0
    l2_correct = 0
    for i, s in enumerate(samples, start=1):
        print(f"  {name} {i}/{len(samples)}")
        result = run_one(s["messages"], debug=debug, l0_threshold=l0_threshold)
        expected = s["correct_label"]
        actual = output_intent(result)
        if actual.get("l1") == expected.get("l1"):
            l1_correct += 1
        if actual.get("l2") == expected.get("l2"):
            l2_correct += 1
        if actual.get("l1") == expected.get("l1") and actual.get("l2") == expected.get("l2"):
            exact += 1
        rows.append([
            messages_text(s["messages"]),
            json.dumps(expected, ensure_ascii=False),
            json.dumps(actual, ensure_ascii=False),
        ])

    total = len(samples)
    acc = exact / total if total else 0.0
    l1_acc = l1_correct / total if total else 0.0
    l2_acc = l2_correct / total if total else 0.0

    wb = Workbook()
    ws = wb.active
    ws.title = name
    ws.append([f"total={total}", f"acc={acc:.4f}", f"L1acc={l1_acc:.4f}", f"L2acc={l2_acc:.4f}"])
    ws.append(["messages", "正确标签", "输出标签（两级意图）"])
    for row in rows:
        ws.append(row)
    ws.column_dimensions["A"].width = 90
    ws.column_dimensions["B"].width = 40
    ws.column_dimensions["C"].width = 40
    path = report_path(f"{name}_report")
    wb.save(path)
    return path


def run_batch(batch: str, debug: bool, l0_threshold: float) -> list[Path]:
    samples = load_samples()
    selected = batch_filter(samples, batch)
    if batch == "batch1":
        return [write_batch1_report(selected, debug=debug, l0_threshold=l0_threshold)]
    if batch == "batch2":
        return [write_risk_report("batch2", selected, debug=debug, l0_threshold=l0_threshold)]
    if batch == "batch3":
        return [write_routing_report("batch3", selected, debug=debug, l0_threshold=l0_threshold)]
    if batch == "batch4":
        return [write_routing_report("batch4", selected, debug=debug, l0_threshold=l0_threshold)]
    if batch == "batch5":
        return [write_risk_report("batch5", selected, debug=debug, l0_threshold=l0_threshold)]
    raise ValueError(batch)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("batch", choices=["batch1", "batch2", "batch3", "batch4", "batch5", "all"])
    parser.add_argument("--debug", action="store_true")
    parser.add_argument("--l0-threshold", type=float, default=0.7)
    args = parser.parse_args(argv)

    batches = ["batch1", "batch2", "batch3", "batch4", "batch5"] if args.batch == "all" else [args.batch]
    paths: list[Path] = []
    for b in batches:
        print(f"Running {b}...")
        paths.extend(run_batch(b, debug=args.debug, l0_threshold=args.l0_threshold))
    for p in paths:
        print(f"Wrote {p}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
