"""
golden_xlsx2json.py — 黄金样本 xlsx → json
==========================================

运营团队用 xlsx 维护黄金样本. 本脚本转换为 prompter 内部 json 格式.

xlsx 格式 (极简):
  Sheet: "黄金样本" (或第一页)
  列A: 用户输入
  列B: 一级意图
  列C: 二级意图

  支持多轮对话: 列A 中用换行分隔各轮, 格式 "角色:内容"
  例:
    用户:我想买保险
    客服:您想了解哪类产品?
    用户:医疗险

用法:
  python golden_xlsx2json.py <golden.xlsx> [--out golden_samples.json]
"""

import json
import sys
from pathlib import Path

from openpyxl import load_workbook


def _find_golden_sheet(wb) -> str | None:
    """查找黄金样本 sheet."""
    # 精确匹配
    for name in ["黄金样本", "golden", "samples", "测试集"]:
        if name in wb.sheetnames:
            return name
    # 兜底: 第一页
    return wb.sheetnames[0] if wb.sheetnames else None


def _parse_conversation(raw: str) -> list[dict]:
    """将列A的文本解析为 messages 列表.

    支持:
      - 单行 → [{"role": "user", "content": "..."}]
      - 多行 (角色:内容) → 逐行解析角色
    """
    raw = raw.strip()
    if not raw:
        return []

    lines = raw.split("\n")
    msgs = []
    for line in lines:
        line = line.strip()
        if not line:
            continue

        # 尝试解析 "角色:内容" 格式
        if ":" in line or "：" in line:
            # 统一冒号
            line = line.replace("：", ":")
            parts = line.split(":", 1)
            role_part = parts[0].strip()
            content = parts[1].strip() if len(parts) > 1 else ""

            if role_part in ("用户", "user"):
                msgs.append({"role": "user", "content": content})
            elif role_part in ("客服", "assistant", "bot"):
                msgs.append({"role": "assistant", "content": content})
            else:
                # 无法识别角色 → 当作 user
                msgs.append({"role": "user", "content": line})
        else:
            msgs.append({"role": "user", "content": line})

    return msgs


def convert(xlsx_path: str | Path, out_path: str | Path | None = None) -> list[dict]:
    """将 ops 黄金样本 xlsx 转为内部 json 格式.

    Returns:
        [{"id": "golden_001", "messages": [...], "expected_l1": "...", "expected_l2": "..."}, ...]
    """
    xlsx_path = Path(xlsx_path)
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    sheet_name = _find_golden_sheet(wb)
    if sheet_name is None:
        raise ValueError(f"No usable sheet found in {xlsx_path.name}")

    ws = wb[sheet_name]
    samples = []

    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=1):
        if not row or len(row) < 3:
            continue

        user_input = str(row[0]).strip() if row[0] else ""
        l1 = str(row[1]).strip() if row[1] else ""
        l2 = str(row[2]).strip() if row[2] else ""

        if not user_input or not l1 or not l2:
            continue

        msgs = _parse_conversation(user_input)
        if not msgs:
            continue

        samples.append({
            "id": f"golden_{i:03d}",
            "messages": msgs,
            "expected_l1": l1,
            "expected_l2": l2,
            "source_batch": "ops",
        })

    wb.close()

    if out_path:
        out_path = Path(out_path)
        out_path.parent.mkdir(parents=True, exist_ok=True)
        with open(out_path, "w", encoding="utf-8") as f:
            json.dump(samples, f, ensure_ascii=False, indent=2)

    return samples


# ═══════════════════════════════════════════════════════════
# CLI
# ═══════════════════════════════════════════════════════════

if __name__ == "__main__":
    import argparse

    p = argparse.ArgumentParser(description="黄金样本 xlsx → json")
    p.add_argument("xlsx", help="黄金样本 xlsx 路径")
    p.add_argument("--out", "-o", default=None, help="输出 json 路径")
    args = p.parse_args()

    samples = convert(args.xlsx, args.out)
    l1_pairs = set((s["expected_l1"], s["expected_l2"]) for s in samples)
    print(f"[OK] {len(samples)} samples, {len(l1_pairs)} unique L1+L2 pairs")
    if args.out:
        print(f"[OK] -> {args.out}")
